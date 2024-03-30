import pandas as pd
from openpyxl import load_workbook

src_file='static/Combined.xlsx'
wb = load_workbook(filename = src_file)
sheet = wb['Data In']
print(sheet.tables.keys())
lookup_table = sheet.tables['TBL_CUR']
data = sheet[lookup_table.ref]
rows_list = []

for row in data:
    cols = []
    for col in row:
        cols.append(col.value)
    rows_list.append(cols)

df = pd.DataFrame(data=rows_list[1:], index=None, columns=rows_list[0])
print(df['CH1'][0])
